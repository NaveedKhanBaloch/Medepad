"""Generated notes as a Word document (.docx), right-to-left for Urdu / Arabic / Persian."""
import io
import re
from datetime import timezone

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.domain import languages
from app.infrastructure.db.models import Generation

TEAL = RGBColor(0x00, 0x79, 0x6B)
_AFTER_BIDI = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
    "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def safe_name(title: str) -> str:
    base = re.sub(r"\s+", " ", re.sub(r"[^\w\- ]+", "", title)).strip()[:80] or "notes"
    return f"{base}.docx"


def _rtl(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.insert_element_before(OxmlElement("w:bidi"), *_AFTER_BIDI)


def note_docx(g: Generation, full: bool = False) -> bytes:
    """The notes; with full=True also the transcript and the exact prompt that produced them."""
    note = g.result or {}
    lang = languages.normalize(g.language) or "en"
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    title = (g.template.name if g.template else "Notes")
    heading = doc.add_heading(title, level=1)
    for run in heading.runs:
        run.font.color.rgb = TEAL
    created = g.created_at.astimezone(timezone.utc) if g.created_at.tzinfo else g.created_at
    version = f"Prompt version {g.prompt.version}" if g.prompt else "Unsaved draft prompt"
    doc.add_paragraph().add_run(f"{version}  ·  {created:%d %b %Y, %H:%M} UTC").italic = True
    if note.get("summary"):
        doc.add_heading(languages.SUMMARY_LABEL.get(lang, "Summary"), level=2)
        doc.add_paragraph(str(note["summary"]))
    for section in note.get("sections", []):
        doc.add_heading(str(section.get("title", "")), level=2)
        for line in str(section.get("content", "")).splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith(("- ", "* ", "• ")):
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(line)
    codes = note.get("diagnostic_codes") or []
    if codes:
        doc.add_heading(languages.CODES_LABEL.get(lang, "Diagnostic codes"), level=2)
        for code in codes:
            doc.add_paragraph(f"{code.get('code', '')} — {code.get('description', '')}", style="List Bullet")
    if not g.success and g.error:
        doc.add_paragraph(f"Generation failed: {g.error}")
    if lang in languages.RTL:
        for paragraph in doc.paragraphs:
            _rtl(paragraph)
    if full:
        doc.add_heading("Transcript", level=1)
        for line in g.transcript.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip())
        doc.add_heading("Prompt used", level=1)
        for line in g.prompt_text.splitlines():
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def results_zip(generations: list[Generation]) -> bytes:
    """Several results in one ZIP: one Word file each, with notes, transcript and prompt."""
    import zipfile

    buf = io.BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for g in generations:
            created = g.created_at.astimezone(timezone.utc) if g.created_at.tzinfo else g.created_at
            version = f"v{g.prompt.version}" if g.prompt else "draft"
            base = safe_name(f"{g.template.name if g.template else 'Free prompt'} {version} {created:%Y-%m-%d %H%M}")[:-5]
            name, n = f"{base}.docx", 2
            while name.lower() in used:
                name, n = f"{base} ({n}).docx", n + 1
            used.add(name.lower())
            z.writestr(name, note_docx(g, full=True))
    return buf.getvalue()
